"""
Generates publication-quality Academic Word (.docx) report for Experiment 5:
Explainable AI (SHAP & LIME) & Algorithmic Fairness Auditing (Fairlearn).
Follows Times New Roman academic styling, incorporates theoretical depth,
structured benchmark tables, and embedded high-resolution figures.
"""

import os
import json
from pathlib import Path
from typing import Dict, Any

import pandas as pd
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

EXPERIMENT_DIR = Path(__file__).resolve().parent.parent
PLOTS_DIR = EXPERIMENT_DIR / "plots"
REPORTS_DIR = EXPERIMENT_DIR / "reports"
MODELS_DIR = EXPERIMENT_DIR / "models"


def set_cell_background(cell, hex_color: str):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    tcPr.append(parse_xml(shd_xml))


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Sets internal padding (in dxa) for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    """Applies subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders_xml = f"""
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))


def generate_experiment_5_docx(output_docx_path: str = None):
    """Generates the comprehensive academic Word report."""
    if output_docx_path is None:
        output_docx_path = str(REPORTS_DIR / "Experiment_5_Report.docx")
    os.makedirs(os.path.dirname(output_docx_path), exist_ok=True)

    summary_file = MODELS_DIR / "exp5_xai_fairness_summary.json"
    benchmark_file = REPORTS_DIR / "bias_mitigation_benchmark.csv"

    metadata = {}
    if summary_file.exists():
        with open(summary_file, "r", encoding="utf-8") as f:
            metadata = json.load(f)

    df_benchmark = pd.read_csv(benchmark_file) if benchmark_file.exists() else pd.DataFrame()

    doc = docx.Document()

    # 0.85-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(3)

    def add_p(text="", bold_prefix=None, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = 'Times New Roman'
            r_bold.font.size = Pt(10.5)
        if text:
            r_text = p.add_run(text)
            r_text.font.name = 'Times New Roman'
            r_text.font.size = Pt(10.5)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(10)
        r_t = p.add_run(text)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(10)
        return p

    def add_callout(text, bold_title="Key Takeaway: "):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_background(cell, "F0F4F8")
        set_cell_margins(cell, top=100, bottom=100, left=160, right=140)

        tcPr = cell._tc.get_or_add_tcPr()
        borders_xml = f"""
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
            <w:left w:val="single" w:sz="24" w:color="1F4E78"/>
        </w:tcBorders>
        """
        tcPr.append(parse_xml(borders_xml))

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r_title = p.add_run(bold_title)
        r_title.bold = True
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        r_text = p.add_run(text)
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def add_image_centered(img_path, width_inches=5.8, caption=None):
        if Path(img_path).exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(str(img_path), width=Inches(width_inches))
            if caption:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(6)
                r = p_cap.add_run(caption)
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # -------------------------------------------------------------------------
    # DOCUMENT HEADER
    # -------------------------------------------------------------------------
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_num.paragraph_format.space_after = Pt(2)
    r_num = p_num.add_run("EXPERIMENT 5 REPORT")
    r_num.bold = True
    r_num.font.name = 'Times New Roman'
    r_num.font.size = Pt(14)
    r_num.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Explainable AI (SHAP & LIME) & Algorithmic Fairness Auditing with Fairlearn")
    r_title.bold = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(0x00, 0x20, 0x60)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(6)
    r_meta = p_meta.add_run("Applied Data Science (ADS)  |  Benchmark: Adult Census Income (N=32,561)  |  September 2026")
    r_meta.italic = True
    r_meta.font.name = 'Times New Roman'
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    add_callout(
        "This academic report investigates Responsible Machine Learning through two intertwined pillars: "
        "model interpretability (SHAP & LIME) and algorithmic fairness (Fairlearn). A champion LightGBM model "
        "(87.69% accuracy, 0.9310 ROC-AUC) is audited for demographic bias and mitigated via pre-processing, "
        "in-processing, and post-processing algorithms.",
        bold_title="Executive Summary: "
    )

    # -------------------------------------------------------------------------
    # SECTION 1: AIM & OBJECTIVES
    # -------------------------------------------------------------------------
    add_heading_1("1. Aim & Objectives")
    add_p("To apply explainable AI (XAI) methods (SHAP & LIME) for interpreting model predictions and evaluate algorithmic fairness using Fairlearn.", bold_prefix="Aim: ")
    add_p("The project fulfills five core research objectives:", bold_prefix="Objectives: ")
    add_bullet("1. Global & Local Model Interpretability: ", "Extract feature importance, directional distributions, and sample-level waterfall attributions using SHAP.")
    add_bullet("2. Local Surrogate Explanations: ", "Deploy LIME tabular explainers across distinct demographic profiles and assess cross-method concordance.")
    add_bullet("3. Fairness Auditing across Protected Groups: ", "Quantify demographic parity, equalized odds, and disparate impact ratios across Sex and Race.")
    add_bullet("4. Tri-Modal Bias Mitigation: ", "Implement and benchmark sample reweighting (pre-processing), ExponentiatedGradient (in-processing), and ThresholdOptimizer (post-processing).")
    add_bullet("5. Empirical Pareto Analysis: ", "Map the quantitative trade-off frontier between model predictive capacity and fairness equity.")

    # -------------------------------------------------------------------------
    # SECTION 2: DATASET & MODEL BENCHMARK
    # -------------------------------------------------------------------------
    add_heading_1("2. Dataset Architecture & Model Selection")
    add_p(
        "We evaluated the gold-standard Adult Census Income dataset (32,561 records, 12 features) predicting whether "
        "an individual's annual income exceeds $50,000. Sensitive demographic attributes include Sex (66.9% Male, 33.1% Female) "
        "and Race (85.4% White, 14.6% Non-White). Following an 80/20 stratified partition, we benchmarked Random Forest "
        "against a champion LightGBM classifier. LightGBM achieved superior discrimination across all evaluation dimensions:"
    )
    add_bullet("Champion LightGBM Classifier: ", "Accuracy: 87.69%, Balanced Accuracy: 80.65%, F1-Score: 0.7240, ROC-AUC: 0.9310.")
    add_bullet("Random Forest Baseline: ", "Accuracy: 86.14%, Balanced Accuracy: 78.34%, F1-Score: 0.6601, ROC-AUC: 0.9203.")

    # -------------------------------------------------------------------------
    # SECTION 3: SHAP INTERPRETABILITY
    # -------------------------------------------------------------------------
    add_heading_1("3. Global & Local Interpretability with SHAP")
    add_p(
        "SHAP grounds model attributions in cooperative game theory, calculating the exact Shapley value for each feature "
        "over all possible feature permutations. This guarantees the four axiomatic properties of efficiency, symmetry, "
        "dummy player, and additivity."
    )
    add_image_centered(PLOTS_DIR / "exp5_shap_summary_bar.png", width_inches=5.6, caption="Figure 1: Global Feature Importance (Mean Absolute SHAP Value Ranking).")
    add_image_centered(PLOTS_DIR / "exp5_shap_beeswarm.png", width_inches=5.8, caption="Figure 2: SHAP Beeswarm Distribution (Feature Value Dispersion vs. Impact on Log-Odds).")
    add_image_centered(PLOTS_DIR / "exp5_shap_dependence.png", width_inches=5.6, caption="Figure 3: SHAP Dependence Plot for Age with Hours per Week Interaction.")
    add_image_centered(PLOTS_DIR / "exp5_shap_waterfall.png", width_inches=5.6, caption="Figure 4: Local SHAP Waterfall Attribution for Test Sample #0.")

    # -------------------------------------------------------------------------
    # SECTION 4: LIME LOCAL EXPLANATIONS & CONCORDANCE
    # -------------------------------------------------------------------------
    add_heading_1("4. Local Explanations with LIME & Cross-XAI Concordance")
    add_p(
        "LIME fits an interpretable sparse linear surrogate model within the localized neighborhood of individual target samples. "
        "We evaluated positive, negative, and borderline prediction profiles, observing clear decision boundaries based on education, "
        "age, and weekly work hours."
    )
    add_image_centered(PLOTS_DIR / "exp5_lime_local_explanations.png", width_inches=6.2, caption="Figure 5: LIME Tabular Explanations Across Positive, Negative, and Borderline Profiles.")
    add_image_centered(PLOTS_DIR / "exp5_xai_shap_vs_lime_comparison.png", width_inches=5.6, caption="Figure 6: Cross-XAI Concordance Audit Comparing SHAP and LIME on Sample #0.")

    # -------------------------------------------------------------------------
    # SECTION 5: FAIRNESS AUDIT & MITIGATION BENCHMARK
    # -------------------------------------------------------------------------
    add_heading_1("5. Algorithmic Fairness Audit & Tri-Modal Mitigation")
    add_p(
        "Fairlearn's MetricFrame revealed severe baseline demographic disparity: the unmitigated model predicted a 26.2% positive rate "
        "for male individuals compared to only 9.1% for females (Demographic Parity Difference = 0.1714; Disparate Impact Ratio = 0.346). "
        "We evaluated three distinct mitigation techniques to restore equity:"
    )

    # Insert Benchmark Table
    if not df_benchmark.empty:
        table = doc.add_table(rows=1, cols=7)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Model Pipeline", "Accuracy", "Bal Acc", "F1", "DP Diff (Sex)", "EO Diff (Sex)", "DIR (Sex)"]
        hdr_cells = table.rows[0].cells
        for idx, h in enumerate(headers):
            hdr_cells[idx].text = h
            set_cell_background(hdr_cells[idx], "1F4E78")
            set_cell_margins(hdr_cells[idx], top=80, bottom=80, left=60, right=60)
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for _, row in df_benchmark.iterrows():
            row_cells = table.add_row().cells
            vals = [
                row["Model Pipeline"].split("(")[0].strip(),
                f"{row['Accuracy']*100:.2f}%",
                f"{row['Balanced Accuracy']*100:.2f}%",
                f"{row['F1-Score']:.4f}",
                f"{row['Demographic Parity Diff (Sex)']:.4f}",
                f"{row['Equalized Odds Diff (Sex)']:.4f}",
                f"{row['Disparate Impact Ratio (Sex)']:.4f}"
            ]
            for c_idx, val in enumerate(vals):
                row_cells[c_idx].text = val
                set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=50, right=50)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].font.size = Pt(8.5)

        set_table_borders(table)
        p_tbl_cap = doc.add_paragraph()
        p_tbl_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_tbl_cap.paragraph_format.space_after = Pt(6)
        r_tc = p_tbl_cap.add_run("Table 1: Quantitative Comparison of Performance vs. Fairness Disparities across Mitigation Strategies.")
        r_tc.italic = True
        r_tc.font.size = Pt(9)

    add_image_centered(PLOTS_DIR / "exp5_fairness_audit_disparity.png", width_inches=5.8, caption="Figure 7: Fairlearn Baseline Demographic Disparities across Gender and Race.")
    add_image_centered(PLOTS_DIR / "exp5_fairness_mitigation_tradeoff.png", width_inches=5.6, caption="Figure 8: Fairness vs. Accuracy Pareto Trade-off Frontier.")

    add_callout(
        "1. Pre-Processing (Sample Reweighting): Produced the most dramatic demographic parity reduction (0.0867 vs 0.1714, a 49.4% drop) and elevated Disparate Impact Ratio to 0.6086.\n"
        "2. Post-Processing (ThresholdOptimizer): Produced the lowest Equalized Odds Difference (0.0393, a 35.6% reduction), calibrating thresholds to equalize error rates across demographic subgroups.",
        bold_title="Mitigation Insights: "
    )

    # -------------------------------------------------------------------------
    # SECTION 6: CONCLUSION
    # -------------------------------------------------------------------------
    add_heading_1("6. Conclusion & Recommendations")
    add_p(
        "Experiment 5 successfully validates that high classification accuracy is insufficient for responsible AI deployment. "
        "By synthesizing game-theoretic SHAP attributions, localized LIME surrogate rules, and Fairlearn demographic auditing, "
        "we comprehensively uncovered systemic bias in socioeconomic modeling. Through pre-processing reweighting and "
        "post-processing threshold optimization, we achieved substantial equity improvements while preserving high operational "
        "performance on the empirical Pareto frontier."
    )

    doc.save(output_docx_path)
    print(f"[+] Academic Word Report generated at: {output_docx_path}")


if __name__ == "__main__":
    generate_experiment_5_docx()
